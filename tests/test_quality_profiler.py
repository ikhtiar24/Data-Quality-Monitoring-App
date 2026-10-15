import tempfile
import unittest
from io import BytesIO
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from quality_profiler import (
    build_exception_frame,
    create_excel_report,
    create_pdf_report,
    evaluate_rules,
    load_history,
    profile_dataframe,
    read_uploaded_file,
    save_run,
    _get_engine,
)


class ProfileTests(unittest.TestCase):
    def test_profile_counts_nulls_duplicates_types_and_distribution(self):
        frame = pd.DataFrame(
            {
                "id": [1, 2, 2, 4],
                "score": [10.0, 101.0, None, 30.0],
                "status": ["ok", "bad", "ok", "ok"],
            }
        )
        profile = profile_dataframe(frame)

        self.assertEqual(profile["row_count"], 4)
        self.assertEqual(profile["column_count"], 3)
        self.assertEqual(profile["duplicate_rows"], 0)
        self.assertEqual(profile["null_cell_count"], 1)
        self.assertEqual(profile["columns"][0]["inferred_type"], "Integer")
        self.assertEqual(profile["columns"][1]["inferred_type"], "Number")
        self.assertTrue(profile["columns"][2]["distribution"])

    def test_profile_counts_duplicate_rows(self):
        frame = pd.DataFrame({"value": ["x", "x", "y"]})
        self.assertEqual(profile_dataframe(frame)["duplicate_rows"], 1)

    def test_rules_report_failures_with_excel_row_numbers(self):
        frame = pd.DataFrame({"id": ["A1", "A1", None], "score": [20, 101, 50]})
        rules = [
            {"column": "id", "expectation": "Unique values"},
            {"column": "id", "expectation": "Never null"},
            {
                "column": "score",
                "expectation": "Between two numbers",
                "minimum": 0,
                "maximum": 100,
            },
            {
                "column": "id",
                "expectation": "Matches a regular expression",
                "pattern": r"^[A-Z]\d$",
            },
            {
                "column": "id",
                "expectation": "Only allowed values",
                "allowed_values": ["A1"],
            },
        ]
        results = evaluate_rules(frame, rules)

        self.assertEqual([result["failed_count"] for result in results], [2, 1, 1, 1, 1])
        self.assertEqual(results[0]["exceptions"][0]["row_number"], 2)
        self.assertEqual(results[2]["exceptions"][0]["row_number"], 3)
        self.assertEqual(len(build_exception_frame(results)), 6)

    def test_csv_upload_and_history_round_trip(self):
        csv = b"customer,amount\nA,5\nB,7\n"
        frame = read_uploaded_file(csv, "customers.csv")
        profile = profile_dataframe(frame)
        rules = [{"column": "amount", "expectation": "Never null"}]
        results = evaluate_rules(frame, rules)
        with tempfile.TemporaryDirectory() as directory:
            database = Path(directory) / "history.sqlite3"
            run_id = save_run("customers.csv", profile, rules, results, database)
            history = load_history(database)
            _get_engine(database).dispose()

        self.assertEqual(len(frame), 2)
        self.assertEqual(int(history.iloc[0]["id"]), run_id)
        self.assertEqual(history.iloc[0]["quality_score"], 100.0)

    def test_xlsx_upload(self):
        output = BytesIO()
        pd.DataFrame({"item": ["A", "B"], "quantity": [2, 5]}).to_excel(
            output, index=False
        )

        frame = read_uploaded_file(output.getvalue(), "items.xlsx")

        self.assertEqual(frame.to_dict(orient="list"), {"item": ["A", "B"], "quantity": [2, 5]})

    def test_unsupported_upload_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "Upload a CSV or Excel"):
            read_uploaded_file(b"data", "data.json")

    def test_excel_and_pdf_reports_are_generated(self):
        frame = pd.DataFrame({"name": ["=1+1", "A"], "score": [110, 90]})
        profile = profile_dataframe(frame)
        rules = [
            {
                "column": "score",
                "expectation": "Between two numbers",
                "minimum": 0,
                "maximum": 100,
            }
        ]
        results = evaluate_rules(frame, rules)

        excel = create_excel_report("demo.csv", profile, results)
        pdf = create_pdf_report("demo.csv", profile, results)
        workbook = load_workbook(BytesIO(excel), data_only=False)

        self.assertTrue(excel.startswith(b"PK"))
        self.assertTrue(pdf.startswith(b"%PDF"))
        self.assertEqual(
            workbook.sheetnames,
            ["Summary", "Column profile", "Rule results", "Exceptions"],
        )


if __name__ == "__main__":
    unittest.main()
